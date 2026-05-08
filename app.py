"""
Compilador Java — Servidor Flask v2
API REST + páginas de resultados
"""

from flask import Flask, render_template, request, jsonify, send_file, session
import json, os , tempfile, io

from src.services.lexer_service import LexerService
from src.services.parser_service import Parser
from src.services.semantic_service import SemanticAnalyzer
from src.services.ast_visualizer import ast_to_text
from src.classes.ast_nodes import *

app = Flask(__name__)
app.secret_key = "compilador_umg_2026"
app.config['MAX_CONTENT_LENGTH'] = 2 * 1024 * 1024

# Cache del último análisis (en memoria)
_last_result = {}


# ================================================================
#  AST → JSON para D3.js
# ================================================================
def ast_to_dict(node):
    COLORS = {
        "ProgramNode": "#2c3e50",      "ClassDeclNode": "#8e44ad",   # Clase: morado
        "MethodDeclNode": "#e91e8c",   "ConstructorDeclNode": "#e91e8c",  # Método: rosa
        "FieldDeclNode": "#2471a3",    "BlockNode": "#7f8c8d",       # Campo: azul
        "VarDeclNode": "#2471a3",      "AssignmentNode": "#d35400",  # Campo/Asig
        "IfNode": "#117a65",           "WhileNode": "#117a65",       # Control: verde esmeralda
        "ForNode": "#117a65",          "ReturnNode": "#117a65",
        "BinaryOpNode": "#00bcd4",     "UnaryOpNode": "#00bcd4",     # Operador: cyan
        "LiteralNode": "#ff6f00",      "IdentifierNode": "#1abc9c",  # Literal: ámbar, Ident: turquesa
        "MethodCallNode": "#f1c40f",   "MemberAccessNode": "#f1c40f",# Llamada: dorado
        "NewObjectNode": "#e91e8c",    "ThisNode": "#e74c3c",
        "ExpressionStmtNode": "#95a5a6", "ParameterNode": "#2471a3",
    }

    def label(n):
        if isinstance(n, ProgramNode): return "Programa"
        if isinstance(n, ClassDeclNode): return n.name
        if isinstance(n, MethodDeclNode): return f"{n.name}()"
        if isinstance(n, ConstructorDeclNode): return f"{n.name}()"
        if isinstance(n, FieldDeclNode): return f"{n.type_name} {n.name}"
        if isinstance(n, BlockNode): return "{ }"
        if isinstance(n, VarDeclNode): return f"{n.type_name} {n.name}"
        if isinstance(n, AssignmentNode): return n.operator
        if isinstance(n, IfNode): return "if"
        if isinstance(n, WhileNode): return "while"
        if isinstance(n, ForNode): return "for"
        if isinstance(n, ReturnNode): return "return"
        if isinstance(n, BinaryOpNode): return n.operator
        if isinstance(n, UnaryOpNode): return n.operator
        if isinstance(n, LiteralNode):
            v = str(n.value)
            return (v[:14] + "..") if len(v) > 16 else v
        if isinstance(n, IdentifierNode): return n.name
        if isinstance(n, MemberAccessNode): return f".{n.member}"
        if isinstance(n, MethodCallNode): return "call()"
        if isinstance(n, NewObjectNode): return f"new {n.class_name}"
        if isinstance(n, ThisNode): return "this"
        if isinstance(n, ExpressionStmtNode): return "expr"
        if isinstance(n, ParameterNode): return f"{n.type_name} {n.name}"
        return type(n).__name__[:10]

    def children(n):
        ch = []
        if isinstance(n, ProgramNode): ch = n.classes
        elif isinstance(n, ClassDeclNode): ch = n.members
        elif isinstance(n, MethodDeclNode): ch = list(n.params) + [n.body]
        elif isinstance(n, ConstructorDeclNode): ch = list(n.params) + [n.body]
        elif isinstance(n, FieldDeclNode):
            if n.initializer: ch = [n.initializer]
        elif isinstance(n, BlockNode): ch = n.statements
        elif isinstance(n, VarDeclNode):
            if n.initializer: ch = [n.initializer]
        elif isinstance(n, AssignmentNode): ch = [n.target, n.value]
        elif isinstance(n, IfNode):
            ch = [n.condition, n.then_block]
            if n.else_block: ch.append(n.else_block)
        elif isinstance(n, WhileNode): ch = [n.condition, n.body]
        elif isinstance(n, ForNode): ch = [n.init, n.condition, n.update, n.body]
        elif isinstance(n, ReturnNode):
            if n.value: ch = [n.value]
        elif isinstance(n, ExpressionStmtNode): ch = [n.expression]
        elif isinstance(n, BinaryOpNode): ch = [n.left, n.right]
        elif isinstance(n, UnaryOpNode): ch = [n.operand]
        elif isinstance(n, MemberAccessNode): ch = [n.object]
        elif isinstance(n, MethodCallNode):
            ch = ([n.callee] if n.callee else []) + list(n.arguments)
        elif isinstance(n, NewObjectNode): ch = list(n.arguments)
        elif isinstance(n, ArrayAccessNode): ch = [n.array, n.index]
        return [c for c in ch if c is not None]

    if node is None: return None
    node_type = type(node).__name__
    return {
        "name": label(node),
        "type": node_type,
        "color": COLORS.get(node_type, "#34495e"),
        "linea": getattr(node, 'line', 0),
        "columna": getattr(node, 'column', 0),
        "children": [c for c in [ast_to_dict(ch) for ch in children(node)] if c]
    }


# ================================================================
#  PÁGINA PRINCIPAL + catch-all para History API
# ================================================================
@app.route("/")
@app.route("/<path:path>")
def index(path=""):
    # Todas las rutas del frontend sirven el mismo index.html
    # El routing real lo hace JavaScript con History API
    # Excluir rutas de API y archivos estáticos
    if path.startswith("api/") or path.startswith("static/"):
        from flask import abort
        abort(404)
    return render_template("index.html")


# ================================================================
#  PÁGINAS DE RESULTADOS (nueva pestaña)
# ================================================================
@app.route("/results/tokens")
def page_tokens():
    return render_template("results/tokens.html",
                           data=_last_result.get("tokens", []),
                           summary=_last_result.get("summary", {}))

@app.route("/results/symbols")
def page_symbols():
    return render_template("results/symbols.html",
                           data=_last_result.get("symbols", []),
                           summary=_last_result.get("summary", {}))

@app.route("/results/ast")
def page_ast():
    return render_template("results/ast.html",
                           ast_text=_last_result.get("ast_text", []),
                           ast_tree=json.dumps(_last_result.get("ast_tree")),
                           summary=_last_result.get("summary", {}))

@app.route("/results/errors/<kind>")
def page_errors(kind):
    key_map = {"lex": "lex_errors", "syn": "syn_errors", "sem": "sem_errors"}
    titles  = {"lex": "Errores Léxicos", "syn": "Errores Sintácticos", "sem": "Errores Semánticos"}
    colors  = {"lex": "#e74c3c", "syn": "#e67e22", "sem": "#d35400"}
    return render_template("results/errors.html",
                           data=_last_result.get(key_map.get(kind, "lex_errors"), []),
                           title=titles.get(kind, "Errores"),
                           color=colors.get(kind, "#e74c3c"),
                           kind=kind,
                           summary=_last_result.get("summary", {}))


# ================================================================
#  API — ANALIZAR
# ================================================================
@app.route("/api/analyze", methods=["POST"])
def analyze():
    global _last_result
    data = request.get_json()
    code = data.get("code", "")
    if not code.strip():
        return jsonify({"error": "No hay código fuente"}), 400

    try:
        lexer  = LexerService(code)
        tokens = lexer.tokenizar()
        tokens_data = [
            {"n": i, "lexema": t.value, "tipo": t.type, "linea": t.line, "columna": t.column}
            for i, t in enumerate(tokens, 1)
        ]
        lex_errors = [
            {"codigo": f"LEX-{i+1:03d}", "lexema": e["lexema"], "linea": e["linea"], "columna": e["columna"], "mensaje": e["descripcion"]}
            for i, e in enumerate(lexer.errores)
        ]
        lex_symbols = [
            {"n": i+1, "nombre": s["nombre"], "tipo": s["tipo"],
             "linea": s["linea"], "columna": s["columna"],
             "valor": s["valor"],
             "ocurrencias": s.get("ocurrencias", 1),
             "lineas": s.get("lineas", [s["linea"]])}
            for i, s in enumerate(lexer.tabla_simbolos)
        ]

        parser = Parser(tokens)
        ast    = parser.parse()
        ast_text = ast_to_text(ast) if ast else []
        ast_json = ast_to_dict(ast)
        syn_errors = [
            {"codigo": e["codigo"], "linea": e["linea"], "columna": e["columna"], "mensaje": e["mensaje"]}
            for e in parser.errors
        ]

        sem = SemanticAnalyzer()
        sem.analyze(ast)
        symbols_data = []
        n = [0]
        def walk(scope, scope_name):
            for name, sym in scope.symbols.items():
                # Omitir built-ins Java (line=0 significa que no vienen del código del usuario)
                if sym.line == 0:
                    continue
                n[0] += 1
                symbols_data.append({
                    "n": n[0], "nombre": sym.name, "categoria": sym.sym_type,
                    "tipo": sym.data_type, "ambito": scope_name,
                    "linea": sym.line, "columna": sym.column
                })
            for child in scope.children: walk(child, child.name)
        walk(sem.symbol_table.global_scope, "global")
        sem_errors = [
            {"codigo": e["codigo"], "linea": e["linea"], "columna": e["columna"], "mensaje": e["mensaje"]}
            for e in sem.errors
        ]

        result = {
            "tokens": tokens_data, "symbols": symbols_data,
            "ast_text": ast_text, "ast_tree": ast_json,
            "lex_errors": lex_errors, "lex_symbols": lex_symbols,
            "syn_errors": syn_errors, "sem_errors": sem_errors,
            "token_stats": lexer.stats,
            "summary": {
                "tokens": len(tokens_data), "symbols": len(symbols_data),
                "lex_errors": len(lex_errors), "lex_symbols": len(lex_symbols),
                "syn_errors": len(syn_errors), "sem_errors": len(sem_errors),
                "identificadores": len(lex_symbols),
            }
        }
        _last_result = result
        return jsonify(result)

    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ================================================================
#  API — EXPORTAR TXT
# ================================================================
@app.route("/api/export/txt")
def export_txt():
    if not _last_result: return "Sin datos", 404
    r = _last_result
    import datetime
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines = ["=" * 70,
             "  COMPILADOR JAVA — REPORTE COMPLETO",
             f"  Universidad Mariano Gálvez de Guatemala",
             f"  Curso: Compiladores  |  Generado: {now}",
             "=" * 70]
    lines += [f"\nTOKENS ({r['summary']['tokens']})", "-" * 70,
              f"{'#':<5} {'TIPO':<22} {'LEXEMA':<20} {'LN':<6} {'COL':<6}", "-" * 70]
    for t in r["tokens"]:
        lex = t["lexema"][:18]
        lines.append(f"{t['n']:<5} {t['tipo']:<22} {lex:<20} {t['linea']:<6} {t['columna']:<6}")
    lines += [f"\nTABLA DE SÍMBOLOS ({r['summary']['symbols']})", "-" * 80,
              f"{'#':<5} {'NOMBRE':<16} {'CAT':<12} {'TIPO':<12} {'ÁMBITO':<18} {'LN':<6}", "-" * 80]
    for s in r["symbols"]:
        lines.append(f"{s['n']:<5} {s['nombre']:<16} {s['categoria']:<12} {s['tipo']:<12} {s['ambito']:<18} {s['linea']:<6}")
    lines += ["\nTABLA LÉXICA DE SÍMBOLOS (" + str(len(r.get("lex_symbols", []))) + ")", "-" * 80,
              f"{'#':<5} {'NOMBRE':<20} {'TIPO':<15} {'OCURR':<8} {'LÍNEAS':<25} {'VALOR':<10}", "-" * 80]
    for s in r.get("lex_symbols", []):
        all_lines = ", ".join(f"Ln {l}" for l in s.get("lineas", [s["linea"]]))
        lines.append(f"{s['n']:<5} {s['nombre']:<20} {s['tipo']:<15} {s.get('ocurrencias',1):<8} {all_lines:<25} {str(s['valor'] or 'N/A'):<10}")
    lines += ["\nÁRBOL AST", "-" * 70] + r["ast_text"]
    for label, key in [("ERRORES LÉXICOS","lex_errors"),("ERRORES SINTÁCTICOS","syn_errors"),("ERRORES SEMÁNTICOS","sem_errors")]:
        errs = r[key]
        lines += [f"\n{label} ({len(errs)})", "-" * 70]
        if errs:
            for e in errs:
                lines.append(f"  [{e['codigo']}] Ln {e['linea']}, Col {e['columna']} | Lexema: '{e.get('lexema','—')}' | {e['mensaje']}")
        else:
            lines.append("  Sin errores.")
    buf = io.BytesIO("\n".join(lines).encode("utf-8"))
    buf.seek(0)
    return send_file(buf, mimetype="text/plain", as_attachment=True, download_name="reporte_compilador.txt")


# ================================================================
#  API — EXPORTAR GRAMÁTICA EBNF como archivo .txt
# ================================================================
@app.route("/api/export/grammar")
def export_grammar():
    grammar = """(* ═══════════════════════════════════════════════════════════ *)
(*  GRAMÁTICA FORMAL — Compilador Java (subconjunto)            *)
(*  Notación: EBNF (Extended Backus-Naur Form)                  *)
(*  Universidad Mariano Gálvez de Guatemala — Compiladores 2026 *)
(* ═══════════════════════════════════════════════════════════ *)

(* PROGRAMA *)
programa            = { declaracion_clase } ;

(* CLASE *)
declaracion_clase   = { modificador } , "class" , IDENTIFICADOR ,
                      [ "extends" , IDENTIFICADOR ] ,
                      [ "implements" , IDENTIFICADOR , { "," , IDENTIFICADOR } ] ,
                      "{" , { miembro_clase } , "}" ;

modificador         = "public" | "private" | "protected"
                    | "static" | "final" | "abstract" ;

miembro_clase       = declaracion_campo
                    | declaracion_metodo
                    | declaracion_constructor ;

(* CAMPOS Y MÉTODOS *)
declaracion_campo   = { modificador } , tipo , IDENTIFICADOR ,
                      [ "=" , expresion ] , ";" ;

declaracion_metodo  = { modificador } , ( tipo | "void" ) , IDENTIFICADOR ,
                      "(" , [ parametros ] , ")" , bloque ;

declaracion_constructor = { modificador } , IDENTIFICADOR ,
                          "(" , [ parametros ] , ")" , bloque ;

parametros          = parametro , { "," , parametro } ;
parametro           = tipo , IDENTIFICADOR ;

(* TIPOS *)
tipo                = tipo_primitivo | IDENTIFICADOR | tipo , "[]" ;
tipo_primitivo      = "int" | "float" | "double" | "boolean"
                    | "char" | "String" | "long" | "short" | "byte" ;

(* SENTENCIAS *)
bloque              = "{" , { sentencia } , "}" ;

sentencia           = declaracion_variable
                    | sentencia_if
                    | sentencia_while
                    | sentencia_for
                    | sentencia_do_while
                    | sentencia_return
                    | sentencia_switch
                    | sentencia_try
                    | sentencia_break
                    | sentencia_continue
                    | sentencia_throw
                    | sentencia_expresion
                    | bloque ;

declaracion_variable = tipo , IDENTIFICADOR , [ "=" , expresion ] , ";" ;

sentencia_if        = "if" , "(" , expresion , ")" , ( bloque | sentencia ) ,
                      [ "else" , ( bloque | sentencia | sentencia_if ) ] ;

sentencia_while     = "while" , "(" , expresion , ")" , ( bloque | sentencia ) ;

sentencia_for       = "for" , "(" ,
                      ( declaracion_variable | sentencia_expresion ) ,
                      expresion , ";" , expresion , ")" ,
                      ( bloque | sentencia ) ;

sentencia_do_while  = "do" , bloque , "while" , "(" , expresion , ")" , ";" ;

sentencia_return    = "return" , [ expresion ] , ";" ;

sentencia_switch    = "switch" , "(" , expresion , ")" , "{" ,
                      { ( "case" , expresion , ":" | "default" , ":" ) ,
                        { sentencia } } , "}" ;

sentencia_try       = "try" , bloque ,
                      { "catch" , "(" , tipo , IDENTIFICADOR , ")" , bloque } ,
                      [ "finally" , bloque ] ;

sentencia_break     = "break" , ";" ;
sentencia_continue  = "continue" , ";" ;
sentencia_throw     = "throw" , expresion , ";" ;

sentencia_expresion = expresion ,
                      [ ( "=" | "+=" | "-=" | "*=" | "/=" | "%=" ) , expresion ] , ";" ;

(* EXPRESIONES — precedencia de menor a mayor *)
expresion           = ternario ;

ternario            = or_logico , [ "?" , or_logico , ":" , or_logico ] ;

or_logico           = and_logico , { "||" , and_logico } ;

and_logico          = igualdad , { "&&" , igualdad } ;

igualdad            = relacional , { ( "==" | "!=" ) , relacional } ;

relacional          = aditiva , { ( "<" | ">" | "<=" | ">=" ) , aditiva } ;

aditiva             = multiplicativa , { ( "+" | "-" ) , multiplicativa } ;

multiplicativa      = unaria , { ( "*" | "/" | "%" ) , unaria } ;

unaria              = ( "!" | "-" | "++" | "--" ) , unaria | postfija ;

postfija            = primaria , [ "++" | "--" ] ;

primaria            = literal
                    | IDENTIFICADOR
                    | "this"
                    | "new" , IDENTIFICADOR , "(" , [ argumentos ] , ")"
                    | "(" , expresion , ")"
                    | acceso_o_llamada ;

acceso_o_llamada    = primaria , { "." , IDENTIFICADOR ,
                      [ "(" , [ argumentos ] , ")" ] } ;

argumentos          = expresion , { "," , expresion } ;

(* LITERALES *)
literal             = INTEGER | FLOAT | STRING | CHAR
                    | "true" | "false" | "null" ;

(* TOKENS TERMINALES *)
IDENTIFICADOR       = letra , { letra | digito | "_" | "$" } ;
INTEGER             = digito , { digito } ;
FLOAT               = digito , { digito } , "." , digito , { digito } ;
STRING              = '"' , { caracter } , '"' ;
CHAR                = "'" , caracter , "'" ;

letra               = "A".."Z" | "a".."z" | "_" | "$" ;
digito              = "0".."9" ;

(* CONVENCIONES EBNF:
   =     definición de regla
   ,     concatenación
   |     alternativa (or)
   { }   cero o más repeticiones
   [ ]   elemento opcional
   " "   token literal terminal
   (* *) comentario
*)"""
    from io import BytesIO
    buf = BytesIO(grammar.encode('utf-8'))
    buf.seek(0)
    return send_file(buf, mimetype="text/plain", as_attachment=True,
                     download_name="gramatica_java_ebnf.txt")


# ================================================================
#  API — EXPORTAR EXCEL
# ================================================================
@app.route("/api/export/excel")
def export_excel():
    if not _last_result: return "Sin datos", 404
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "pip install openpyxl", 500

    r = _last_result
    wb = openpyxl.Workbook(); wb.remove(wb.active)

    def hs(c):
        return dict(font=Font(bold=True,color="FFFFFF",size=11),
                    fill=PatternFill("solid",fgColor=c),
                    alignment=Alignment(horizontal="center",vertical="center"))
    def ap(cell, s):
        for k, v in s.items(): setattr(cell, k, v)
    def aw(ws):
        for col in ws.columns:
            ml = max((len(str(c.value)) for c in col if c.value), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(ml+4, 65)

    FA = PatternFill("solid", fgColor="F2F2F2")
    FB = PatternFill("solid", fgColor="FFFFFF")
    CA = Alignment(horizontal="center")

    def add(title, headers, rows, hcolor, cc=()):
        ws = wb.create_sheet(title=title); ws.freeze_panes = "A2"
        h = hs(hcolor)
        for ci, hd in enumerate(headers, 1): ap(ws.cell(row=1,column=ci,value=hd), h)
        for ri, row in enumerate(rows, 2):
            f = FA if ri%2==0 else FB
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri,column=ci,value=val); c.fill=f
                if ci in cc: c.alignment=CA
        aw(ws)

    add("Tokens", ["#","Lexema","Token","Línea","Columna"],
        [(t["n"],t["lexema"],t["tipo"],t["linea"],t["columna"]) for t in r["tokens"]],
        "2980b9", (1,4,5))
    add("Tabla Léxica", ["#","Nombre","Tipo Token","Ocurrencias","Líneas"],
        [(s["n"],s["nombre"],s["tipo"],s.get("ocurrencias",1),
          ", ".join(f"Ln {l}" for l in s.get("lineas",[s["linea"]]))) for s in r.get("lex_symbols",[])],
        "27ae60", (1,4))
    add("Tabla de Símbolos", ["#","Nombre","Categoría","Tipo","Ámbito","Línea","Col"],
        [(s["n"],s["nombre"],s["categoria"],s["tipo"],s["ambito"],s["linea"],s["columna"]) for s in r["symbols"]],
        "8e44ad", (1,6,7))
    ws = wb.create_sheet(title="Árbol AST")
    ap(ws.cell(row=1,column=1,value="Árbol AST — Representación textual"), hs("27ae60"))
    ws.column_dimensions["A"].width = 80
    for ri, ln in enumerate(r["ast_text"], 2):
        c = ws.cell(row=ri,column=1,value=ln); c.fill=FA if ri%2==0 else FB
        c.font = Font(name="Consolas",size=10)

    # Imagen del árbol generada con PIL puro (sin tkinter)
    img_row = len(r["ast_text"]) + 4
    tmp_path = None
    try:
        from PIL import Image as PILImage, ImageDraw
        import openpyxl.drawing.image as xl_img

        ast_tree = r.get("ast_tree")
        if ast_tree:
            NODE_R = 28
            H_GAP  = 16
            V_GAP  = 80
            COLORS_MAP = {
                "ProgramNode":"#2c3e50",    "ClassDeclNode":"#8e44ad",
                "MethodDeclNode":"#e91e8c", "ConstructorDeclNode":"#e91e8c",
                "FieldDeclNode":"#2471a3",  "BlockNode":"#7f8c8d",
                "VarDeclNode":"#2471a3",    "AssignmentNode":"#d35400",
                "IfNode":"#117a65",         "WhileNode":"#117a65","ForNode":"#117a65",
                "ReturnNode":"#117a65",     "BinaryOpNode":"#00bcd4","UnaryOpNode":"#00bcd4",
                "LiteralNode":"#ff6f00",    "IdentifierNode":"#1abc9c",
                "MethodCallNode":"#f1c40f", "MemberAccessNode":"#f1c40f",
                "NewObjectNode":"#e91e8c",  "ThisNode":"#e74c3c",
                "ExpressionStmtNode":"#95a5a6","ParameterNode":"#2471a3",
            }
            DEFAULT_C = "#34495e"

            def build(n, depth, x_off):
                label    = n.get("name","?")
                ntype    = n.get("type","")
                kids     = [build(c, depth+1, 0) for c in n.get("children",[])]
                return {"label":label,"type":ntype,"children":kids,"x":0,"y":depth*V_GAP+50}

            def calc_pos(t, depth, x_off):
                t["y"] = depth * V_GAP + 50
                if not t["children"]:
                    t["x"] = x_off + NODE_R
                    return x_off + NODE_R*2 + H_GAP
                cx = x_off
                for ch in t["children"]:
                    cx = calc_pos(ch, depth+1, cx)
                t["x"] = (t["children"][0]["x"] + t["children"][-1]["x"]) / 2
                return max(cx, t["x"] + NODE_R + H_GAP)

            def shift(t, dx):
                t["x"] += dx
                for ch in t["children"]: shift(ch, dx)

            def find_max(t, key):
                return max(t[key], *(find_max(c,key) for c in t["children"]) if t["children"] else [t[key]])
            def find_min_x(t):
                return min(t["x"], *(find_min_x(c) for c in t["children"]) if t["children"] else [t["x"]])

            td = build(ast_tree, 0, 0)
            calc_pos(td, 0, 0)
            shift(td, -find_min_x(td) + 60)

            max_x = int(find_max(td,"x")) + 120
            max_y = int(find_max(td,"y")) + 120
            img   = PILImage.new("RGB", (max(max_x,400), max(max_y,300)), "white")
            draw  = ImageDraw.Draw(img)

            def hex_rgb(h):
                h = h.lstrip("#")
                return tuple(int(h[i:i+2],16) for i in (0,2,4))

            def draw_edges(t):
                for ch in t["children"]:
                    draw.line([(int(t["x"]),int(t["y"])+NODE_R),(int(ch["x"]),int(ch["y"])-NODE_R)],
                              fill=(127,140,141), width=2)
                    draw_edges(ch)

            def draw_nodes(t):
                x,y,label = int(t["x"]),int(t["y"]),t["label"]
                bg = hex_rgb(COLORS_MAP.get(t["type"], DEFAULT_C))
                hw = max(NODE_R, len(label)*7.5/2+10)
                draw.ellipse([x-hw,y-NODE_R,x+hw,y+NODE_R], fill=bg, outline=(44,62,80), width=2)
                try:
                    bbox = draw.textbbox((0,0), label)
                    tw, th = bbox[2]-bbox[0], bbox[3]-bbox[1]
                except AttributeError:
                    tw, th = draw.textsize(label)
                draw.text((x - tw//2, y - th//2), label, fill=(255,255,255))
                for ch in t["children"]: draw_nodes(ch)

            draw_edges(td); draw_nodes(td)

            import tempfile, os
            tmp_path = os.path.join(tempfile.gettempdir(), "ast_web_export.png")
            img.save(tmp_path, format="PNG")

            xl_image = xl_img.Image(tmp_path)
            MAX_W = 900
            if xl_image.width > MAX_W:
                f = MAX_W / xl_image.width
                xl_image.width  = int(xl_image.width  * f)
                xl_image.height = int(xl_image.height * f)

            ws.cell(row=img_row, column=1, value="── Imagen del Árbol AST ──")
            ws.add_image(xl_image, f"A{img_row+1}")
        else:
            ws.cell(row=img_row, column=1, value="(Analiza código primero para incluir imagen)")
    except ImportError:
        ws.cell(row=img_row, column=1, value="(Instala pillow: pip install pillow)")
    except Exception as ex:
        ws.cell(row=img_row, column=1, value=f"(Error imagen: {ex})")
    for label, key, color in [("Errores Léxicos","lex_errors","c0392b"),
                               ("Errores Sintácticos","syn_errors","e67e22"),
                               ("Errores Semánticos","sem_errors","d35400")]:
        errs = r[key]
        if key == "lex_errors":
            rows = [(e["codigo"],e.get("lexema","—"),e["linea"],e["columna"],e["mensaje"]) for e in errs] or [("—","—","—","—","Sin errores")]
            add(label, ["Código","Lexema","Línea","Col","Descripción"], rows, color, (1,3,4))
        else:
            rows = [(e["codigo"],e["linea"],e["columna"],e["mensaje"]) for e in errs] or [("—","—","—","Sin errores")]
            add(label, ["Código","Línea","Col","Descripción"], rows, color, (1,2,3))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    try:
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)
    except Exception:
        pass
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="reporte_compilador.xlsx")


if __name__ == "__main__":
    print("\n  Compilador Java — Servidor Web v2")
    print("  Abre http://localhost:5000 en tu navegador\n")
    app.run(debug=True, port=5000)
